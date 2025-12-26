#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
============================================
Project: FormFlow
File: excel_handlers.py
Created: 2025-12-16
Author: @lewopxd

Description:
Excel file handling with efficient caching system.
- Fast initial load (metadata + active sheet only)
- Background parsing of all sheets
- In-memory cache for instant sheet switching
- Support for restore from project file
============================================
"""

from pathlib import Path
from typing import Dict, Any, List, Optional
import webview
import threading
from datetime import datetime

# Import Logger
try:
    from core.logger import Logger
except ImportError:
    class Logger:
        @staticmethod
        def debug(msg): print(f"DEBUG: {msg}")
        @staticmethod
        def info(msg): print(f"INFO: {msg}")
        @staticmethod
        def error(msg): print(f"ERROR: {msg}")


class ExcelCache:
    """In-memory cache for parsed Excel data"""
    
    def __init__(self):
        self.path: Optional[str] = None
        self.filename: Optional[str] = None
        self.sheets: List[str] = []
        self.active_sheet: Optional[str] = None
        self.cached_data: Dict[str, Dict] = {}  # {sheet_name: {headers, rows}}
        self.is_fully_parsed: bool = False
        self.lock = threading.Lock()
    
    def clear(self):
        """Clear all cached data"""
        with self.lock:
            self.path = None
            self.filename = None
            self.sheets = []
            self.active_sheet = None
            self.cached_data = {}
            self.is_fully_parsed = False
    
    def set_metadata(self, path: str, filename: str, sheets: List[str], active_sheet: str):
        """Set file metadata"""
        with self.lock:
            self.path = path
            self.filename = filename
            self.sheets = sheets
            self.active_sheet = active_sheet
    
    def cache_sheet(self, sheet_name: str, headers: List[Dict], rows: List[List]):
        """Cache a single sheet's data"""
        with self.lock:
            self.cached_data[sheet_name] = {
                "headers": headers,
                "rows": rows,
                "rowCount": len(rows),
                "colCount": len(headers)
            }
    
    def get_sheet(self, sheet_name: str) -> Optional[Dict]:
        """Get cached sheet data"""
        with self.lock:
            return self.cached_data.get(sheet_name)
    
    def has_sheet(self, sheet_name: str) -> bool:
        """Check if sheet is cached"""
        with self.lock:
            return sheet_name in self.cached_data
    
    def to_dict(self) -> Dict:
        """Export cache for project save"""
        with self.lock:
            return {
                "path": self.path,
                "filename": self.filename,
                "sheets": self.sheets,
                "activeSheet": self.active_sheet,
                "cachedData": self.cached_data.copy(),
                "isFullyParsed": self.is_fully_parsed
            }
    
    def from_dict(self, data: Dict):
        """Restore cache from project data"""
        with self.lock:
            self.path = data.get("path")
            self.filename = data.get("filename")
            self.sheets = data.get("sheets", [])
            self.active_sheet = data.get("activeSheet")
            self.cached_data = data.get("cachedData", {})
            self.is_fully_parsed = data.get("isFullyParsed", False)


# Global cache instance
excel_cache = ExcelCache()


class ExcelHandler:
    def __init__(self, bridge):
        self.bridge = bridge
        self.current_wb = None

    def register(self):
        """Register handlers with the bridge"""
        self.bridge.register_handler("excel_browse", self.handle_browse)
        self.bridge.register_handler("excel_parse", self.handle_parse)
        self.bridge.register_handler("excel_browse_and_parse", self.handle_browse_and_parse)
        self.bridge.register_handler("excel_switch_sheet", self.handle_switch_sheet)
        self.bridge.register_handler("excel_reload", self.handle_reload)
        self.bridge.register_handler("excel_get_cache", self.handle_get_cache)
        self.bridge.register_handler("excel_restore_cache", self.handle_restore_cache)
        # Export handlers
        self.bridge.register_handler("get_default_export_path", self.handle_get_default_export_path)
        self.bridge.register_handler("browse_export_folder", self.handle_browse_export_folder)
        self.bridge.register_handler("export_form_responses", self.handle_export_form_responses)
        self.bridge.register_handler("export_automation_package", self.handle_export_automation_package)
        
        # Warmup openpyxl in background thread
        threading.Thread(target=self._warmup_openpyxl, daemon=True).start()

    def _warmup_openpyxl(self):
        """Pre-load openpyxl module to speed up first user interaction"""
        try:
            Logger.info("[Excel] Warming up openpyxl...")
            import openpyxl
            from openpyxl import Workbook
            wb = Workbook()
            _ = wb.active
            _ = openpyxl.load_workbook
            Logger.info("[Excel] Warmup complete")
        except Exception as e:
            Logger.error(f"[Excel] Warmup failed: {e}")

    def get_column_letter(self, index: int) -> str:
        """Convert column index to Excel letter (0=A, 1=B, 26=AA, etc.)"""
        result = ""
        while index >= 0:
            result = chr(65 + index % 26) + result
            index = index // 26 - 1
        return result

    def _serialize_cell_value(self, value) -> str:
        """Convert cell value to string, handling dates, numbers, etc."""
        from datetime import datetime, date, time, timedelta
        from decimal import Decimal
        
        if value is None:
            return ""
        
        if isinstance(value, datetime):
            if value.hour or value.minute or value.second:
                return value.strftime("%Y-%m-%d %H:%M")
            return value.strftime("%Y-%m-%d")
        
        if isinstance(value, date):
            return value.strftime("%Y-%m-%d")
        
        if isinstance(value, time):
            return value.strftime("%H:%M:%S")
        
        if isinstance(value, timedelta):
            return str(value.total_seconds())
        
        if isinstance(value, Decimal):
            return str(float(value))
        
        if isinstance(value, bytes):
            return ""
        
        if isinstance(value, float):
            if value == int(value):
                return str(int(value))
            return str(value)
        
        return str(value).strip()

    def _parse_sheet(self, wb, sheet_name: str) -> Dict:
        """Parse a single sheet and return headers + rows"""
        sheet = wb[sheet_name]
        
        raw_data = []
        for row in sheet.iter_rows(values_only=True):
            row_list = [self._serialize_cell_value(cell) for cell in row]
            raw_data.append(row_list)
        
        # Trim trailing empty rows
        while raw_data and not any(cell.strip() for cell in raw_data[-1]):
            raw_data.pop()
        
        # Extract headers from first row
        headers = []
        if raw_data:
            header_row = raw_data[0]
            for i, val in enumerate(header_row):
                headers.append({
                    "index": i,
                    "letter": self.get_column_letter(i),
                    "name": str(val).strip() if val else f"Col{i+1}"
                })
            rows = raw_data[1:]
        else:
            rows = []
        
        return {"headers": headers, "rows": rows}

    def handle_browse(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Open native file dialog"""
        try:
            windows = webview.windows
            if not windows:
                return {"error": "No active window", "cancelled": True}
            
            window = windows[0]
            result = window.create_file_dialog(
                webview.OPEN_DIALOG,
                allow_multiple=False,
                file_types=('Excel Files (*.xlsx;*.xls;*.xlsm)', 'All Files (*.*)')
            )
            
            if result and len(result) > 0:
                filepath = result[0]
                Logger.info(f"[Excel] Selected: {Path(filepath).name}")
                return {"path": filepath, "cancelled": False}
            
            return {"path": None, "cancelled": True}
            
        except Exception as e:
            Logger.error(f"[Excel] Browse error: {e}")
            return {"error": str(e), "cancelled": True}

    def handle_parse(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """
        Parse Excel file with efficient caching.
        Flow:
        1. Send metadata immediately
        2. Parse and send active sheet
        3. Background parse all other sheets
        """
        filepath = content.get("path", "")
        sheet_name = content.get("sheet", None)
        
        if not filepath:
            return {"error": "No path provided"}
        
        # Clear existing cache
        excel_cache.clear()
        
        # Start parsing in thread
        threading.Thread(
            target=self._process_excel_file, 
            args=(filepath, sheet_name),
            daemon=True
        ).start()
        
        return {"status": "processing", "message": "Started parsing"}

    def handle_browse_and_parse(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Combined browse and parse"""
        browse_result = self.handle_browse(content)
        
        if browse_result.get("cancelled") or not browse_result.get("path"):
            return {"cancelled": True}
        
        self.handle_parse({"path": browse_result["path"]})
        return {"status": "processing", "path": browse_result["path"]}

    def handle_switch_sheet(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Switch to a different sheet - uses cache, no re-parsing"""
        sheet_name = content.get("sheet")
        
        if not sheet_name:
            return {"error": "No sheet name provided"}
        
        if not excel_cache.has_sheet(sheet_name):
            # Sheet not cached yet, need to parse it
            Logger.info(f"[Excel] Sheet '{sheet_name}' not cached, parsing...")
            return self._parse_and_send_sheet(sheet_name)
        
        # Get from cache
        cached = excel_cache.get_sheet(sheet_name)
        if not cached:
            return {"error": f"Sheet '{sheet_name}' not found in cache"}
        
        # Update active sheet
        excel_cache.active_sheet = sheet_name
        
        # Send cached data
        data_payload = {
            "filename": excel_cache.filename,
            "activeSheet": sheet_name,
            "sheets": excel_cache.sheets,
            "headers": cached["headers"],
            "data": cached["rows"],
            "rowCount": cached["rowCount"],
            "colCount": cached["colCount"],
            "fromCache": True
        }
        self.bridge.send_to_js("excel_data", data_payload)
        Logger.info(f"[Excel] Sent cached data for '{sheet_name}'")
        
        return {"success": True, "fromCache": True}

    def handle_reload(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Reload Excel file from disk - re-parse everything"""
        path = excel_cache.path
        
        if not path:
            return {"error": "No file to reload"}
        
        if not Path(path).exists():
            return {"error": f"File not found: {path}"}
        
        Logger.info(f"[Excel] Reloading file: {path}")
        
        # Clear cache and re-parse
        active = excel_cache.active_sheet
        excel_cache.clear()
        
        threading.Thread(
            target=self._process_excel_file, 
            args=(path, active),
            daemon=True
        ).start()
        
        return {"status": "reloading"}

    def handle_get_cache(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Get current cache for project save"""
        return excel_cache.to_dict()

    def handle_restore_cache(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Restore cache from project data and send to UI"""
        excel_data = content.get("excel")
        
        if not excel_data:
            return {"success": False, "error": "No excel data provided"}
        
        try:
            # Restore cache from project data
            excel_cache.from_dict(excel_data)
            Logger.info(f"[Excel] Restored cache: {excel_cache.filename}, {len(excel_cache.sheets)} sheets")
            
            # Send metadata to UI
            meta_payload = {
                "path": excel_cache.path,
                "filename": excel_cache.filename,
                "sheets": excel_cache.sheets,
                "activeSheet": excel_cache.active_sheet,
                "sizeBytes": 0,  # Not available from cache
                "fromCache": True
            }
            self.bridge.send_to_js("excel_meta", meta_payload)
            
            # Send active sheet data
            active_sheet = excel_cache.active_sheet
            if active_sheet and excel_cache.has_sheet(active_sheet):
                cached = excel_cache.get_sheet(active_sheet)
                data_payload = {
                    "filename": excel_cache.filename,
                    "activeSheet": active_sheet,
                    "sheets": excel_cache.sheets,
                    "headers": cached["headers"],
                    "data": cached["rows"],
                    "rowCount": cached["rowCount"],
                    "colCount": cached["colCount"],
                    "fromCache": True
                }
                self.bridge.send_to_js("excel_data", data_payload)
                Logger.info(f"[Excel] Sent restored data for '{active_sheet}'")
            
            return {"success": True}
            
        except Exception as e:
            Logger.error(f"[Excel] Restore cache error: {e}")
            return {"success": False, "error": str(e)}

    def _process_excel_file(self, filepath: str, sheet_name: str = None):
        """Worker method to parse Excel with caching"""
        try:
            import openpyxl
            path = Path(filepath)
            
            Logger.info(f"[Excel] Processing: {path.name}")
            
            # Open workbook
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            
            # Determine active sheet
            if sheet_name and sheet_name in sheet_names:
                active_sheet = sheet_name
            else:
                active_sheet = wb.active.title
            
            # Set metadata in cache
            excel_cache.set_metadata(
                path=str(path.absolute()),
                filename=path.name,
                sheets=sheet_names,
                active_sheet=active_sheet
            )
            
            # Send Metadata Event immediately
            meta_payload = {
                "path": str(path.absolute()),
                "filename": path.name,
                "sheets": sheet_names,
                "activeSheet": active_sheet,
                "sizeBytes": path.stat().st_size
            }
            self.bridge.send_to_js("excel_meta", meta_payload)
            Logger.info("[Excel] Sent metadata")
            
            # Parse active sheet first
            Logger.info(f"[Excel] Parsing active sheet: {active_sheet}")
            parsed = self._parse_sheet(wb, active_sheet)
            excel_cache.cache_sheet(active_sheet, parsed["headers"], parsed["rows"])
            
            # Send active sheet data
            data_payload = {
                "filename": path.name,
                "activeSheet": active_sheet,
                "sheets": sheet_names,
                "headers": parsed["headers"],
                "data": parsed["rows"],
                "rowCount": len(parsed["rows"]),
                "colCount": len(parsed["headers"])
            }
            self.bridge.send_to_js("excel_data", data_payload)
            Logger.info(f"[Excel] Sent active sheet: {len(parsed['rows'])} rows")
            
            # Background parse all other sheets
            for sname in sheet_names:
                if sname != active_sheet:
                    Logger.debug(f"[Excel] Background parsing: {sname}")
                    parsed = self._parse_sheet(wb, sname)
                    excel_cache.cache_sheet(sname, parsed["headers"], parsed["rows"])
            
            excel_cache.is_fully_parsed = True
            wb.close()
            
            Logger.info(f"[Excel] All {len(sheet_names)} sheets cached")
            
            # Notify UI that background parsing is complete
            self.bridge.send_to_js("excel_cache_ready", {
                "sheets": sheet_names,
                "activeSheet": active_sheet
            })
            
        except Exception as e:
            Logger.error(f"[Excel] Processing error: {e}")
            import traceback
            Logger.debug(traceback.format_exc())
            self.bridge.send_to_js("excel_error", {"error": str(e)})

    def _parse_and_send_sheet(self, sheet_name: str) -> Dict:
        """Parse a single sheet on-demand (fallback if not cached)"""
        if not excel_cache.path:
            return {"error": "No file loaded"}
        
        try:
            import openpyxl
            wb = openpyxl.load_workbook(excel_cache.path, read_only=True, data_only=True)
            
            if sheet_name not in wb.sheetnames:
                wb.close()
                return {"error": f"Sheet '{sheet_name}' not found"}
            
            parsed = self._parse_sheet(wb, sheet_name)
            excel_cache.cache_sheet(sheet_name, parsed["headers"], parsed["rows"])
            excel_cache.active_sheet = sheet_name
            
            wb.close()
            
            # Send data
            data_payload = {
                "filename": excel_cache.filename,
                "activeSheet": sheet_name,
                "sheets": excel_cache.sheets,
                "headers": parsed["headers"],
                "data": parsed["rows"],
                "rowCount": len(parsed["rows"]),
                "colCount": len(parsed["headers"])
            }
            self.bridge.send_to_js("excel_data", data_payload)
            
            return {"success": True}
            
        except Exception as e:
            Logger.error(f"[Excel] Parse sheet error: {e}")
            return {"error": str(e)}

    def handle_get_default_export_path(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Get default export path (Documents/autoforms_exports/)"""  
        try:
            import os
            # Get My Documents folder
            if os.name == 'nt':  # Windows
                documents_path = Path.home() / "Documents"
            else:  # Linux/Mac
                documents_path = Path.home() / "Documents"
            
            # Default export subfolder
            export_folder = documents_path / "autoforms_exports"
            
            # Return path (don't create yet - only create on actual export)
            Logger.info(f"[Excel Export] Default path: {export_folder}")
            return {"success": True, "path": str(export_folder)}
            
        except Exception as e:
            Logger.error(f"[Excel Export] Error getting default path: {e}")
            return {"success": False, "error": str(e)}

    def handle_browse_export_folder(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """Open native folder selection dialog for export destination"""
        try:
            windows = webview.windows
            if not windows:
                return {"success": False, "error": "No active window"}
            
            window = windows[0]
            result = window.create_file_dialog(
                webview.FOLDER_DIALOG,
                allow_multiple=False
            )
            
            if result and len(result) > 0:
                folder_path = result[0]
                Logger.info(f"[Excel Export] Selected folder: {folder_path}")
                return {"success": True, "path": folder_path}
            
            return {"success": False, "cancelled": True}
            
        except Exception as e:
            Logger.error(f"[Excel Export] Browse folder error: {e}")
            return {"success": False, "error": str(e)}

    def handle_export_form_responses(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """
        Export form responses to Excel file.
        
        Expected content:
        - filename: str - Name of output file (with .xlsx)
        - output_path: str - Directory path for output
        - headers: List[str] - Column headers
        - rows: List[List[str]] - Row data
        - options: Dict with Excel options (sheetName, createAsTable, tableName, wrapText, freezePanes, autoWidth, columnWidth)
        """
        try:
            import openpyxl
            import re
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.worksheet.table import Table, TableStyleInfo
            
            filename = content.get("filename", "export.xlsx")
            output_path = content.get("output_path", "")
            headers = content.get("headers", [])
            rows = content.get("rows", [])
            
            # Get Excel options with defaults
            options = content.get("options", {})
            sheet_name = options.get("sheetName", "Respuestas")
            create_as_table = options.get("createAsTable", True)
            table_name = options.get("tableName", "Table_1")
            wrap_text = options.get("wrapText", False)  # Default: no wrap (truncate)
            freeze_panes = options.get("freezePanes", True)
            auto_width = options.get("autoWidth", True)
            column_width = options.get("columnWidth", 50)
            
            if not output_path:
                return {"success": False, "error": "No output path provided"}
            
            if not headers:
                return {"success": False, "error": "No headers provided"}
            
            # Create output folder if it doesn't exist
            output_dir = Path(output_path)
            if not output_dir.exists():
                output_dir.mkdir(parents=True, exist_ok=True)
                Logger.info(f"[Excel Export] Created folder: {output_dir}")
            
            # Create workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name[:31]  # Excel sheet name max 31 chars
            
            # Style definitions
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell_alignment = Alignment(vertical="center", wrap_text=wrap_text)
            thin_border = Border(
                left=Side(style='thin', color='E5E7EB'),
                right=Side(style='thin', color='E5E7EB'),
                top=Side(style='thin', color='E5E7EB'),
                bottom=Side(style='thin', color='E5E7EB')
            )
            
            # Write headers
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # Write data rows
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
                    cell.alignment = cell_alignment
            
            # Auto-adjust column widths
            for col_idx, header in enumerate(headers, 1):
                if auto_width:
                    max_length = len(str(header))
                    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                        for cell in row:
                            try:
                                if cell.value:
                                    # For auto-width, consider first line only if wrap is off
                                    value_str = str(cell.value)
                                    if not wrap_text:
                                        value_str = value_str.split('\n')[0]
                                    max_length = max(max_length, min(len(value_str), column_width))
                            except:
                                pass
                    adjusted_width = min(max_length + 2, column_width)
                else:
                    adjusted_width = column_width
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = adjusted_width
            
            # Create as Excel Table
            if create_as_table and rows:
                last_col = openpyxl.utils.get_column_letter(len(headers))
                last_row = len(rows) + 1  # +1 for header
                table_ref = f"A1:{last_col}{last_row}"
                
                # Sanitize table name (no spaces, must start with letter/underscore)
                safe_table_name = re.sub(r'[^a-zA-Z0-9_]', '_', table_name)
                if safe_table_name and not safe_table_name[0].isalpha() and safe_table_name[0] != '_':
                    safe_table_name = 'T_' + safe_table_name
                if not safe_table_name:
                    safe_table_name = 'Table_1'
                
                table = Table(displayName=safe_table_name, ref=table_ref)
                style = TableStyleInfo(
                    name="TableStyleMedium9",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False
                )
                table.tableStyleInfo = style
                ws.add_table(table)
            
            # Freeze header row
            if freeze_panes:
                ws.freeze_panes = 'A2'
            
            # Build full path
            full_path = output_dir / filename
            
            # Save workbook
            wb.save(full_path)
            wb.close()
            
            Logger.info(f"[Excel Export] Saved: {full_path} ({len(rows)} rows)")
            return {"success": True, "path": str(full_path), "rowCount": len(rows)}
            
        except Exception as e:
            Logger.error(f"[Excel Export] Error: {e}")
            import traceback
            Logger.debug(traceback.format_exc())
            return {"success": False, "error": str(e)}

    def handle_export_automation_package(self, content: Dict[str, Any]) -> Dict[str, Any]:
        """
        Export automation package (.afpkg) as JSON file.
        
        Expected content:
        - filename: str - Name of output file (with .afpkg)
        - output_path: str - Directory path for output
        - package: Dict - The complete automation package (meta, instructions, resolvedRows)
        """
        try:
            import json
            
            filename = content.get("filename", "automation.afpkg")
            output_path = content.get("output_path", "")
            package = content.get("package", {})
            
            if not output_path:
                return {"success": False, "error": "No output path provided"}
            
            if not package:
                return {"success": False, "error": "No package data provided"}
            
            # Create output folder if it doesn't exist
            output_dir = Path(output_path)
            if not output_dir.exists():
                output_dir.mkdir(parents=True, exist_ok=True)
                Logger.info(f"[Automation Export] Created folder: {output_dir}")
            
            # Build full path
            full_path = output_dir / filename
            
            # Save as JSON with formatting
            with open(full_path, 'w', encoding='utf-8') as f:
                json.dump(package, f, ensure_ascii=False, indent=2)
            
            row_count = len(package.get("resolvedRows", []))
            Logger.info(f"[Automation Export] Saved: {full_path} ({row_count} rows)")
            return {"success": True, "path": str(full_path), "rowCount": row_count}
            
        except Exception as e:
            Logger.error(f"[Automation Export] Error: {e}")
            import traceback
            Logger.debug(traceback.format_exc())
            return {"success": False, "error": str(e)}


# Factory function
def register_excel_handlers(bridge):
    handler = ExcelHandler(bridge)
    handler.register()
    return handler  # Return for access to cache
