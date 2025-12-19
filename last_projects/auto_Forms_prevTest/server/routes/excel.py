"""
Excel File Handling Routes
Parse Excel files using openpyxl on backend
"""
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel
from pathlib import Path
from typing import Optional, List
import os

router = APIRouter(prefix="/api/excel", tags=["excel"])


class ExcelParseRequest(BaseModel):
    path: str
    sheet: Optional[str] = None


class ColumnInfo(BaseModel):
    index: int
    letter: str
    name: str


@router.post("/parse")
async def parse_excel(body: ExcelParseRequest):
    """
    Parse Excel file and return data.
    Uses openpyxl for efficient backend parsing.
    """
    path = Path(body.path)
    
    if not path.exists():
        raise HTTPException(404, f"File not found: {body.path}")
    
    if not path.suffix.lower() in ['.xlsx', '.xls', '.xlsm']:
        raise HTTPException(400, f"Invalid file type: {path.suffix}")
    
    try:
        import openpyxl
        
        # Load workbook (read_only for performance, data_only for calculated values)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        
        sheet_names = wb.sheetnames
        
        # Select sheet
        if body.sheet and body.sheet in sheet_names:
            sheet = wb[body.sheet]
        else:
            sheet = wb.active
        
        # Read all data
        data = []
        for row in sheet.iter_rows(values_only=True):
            # Convert all cells to strings, handle None
            row_data = []
            for cell in row:
                if cell is None:
                    row_data.append("")
                elif isinstance(cell, (int, float)):
                    # Keep numbers as-is for proper formatting
                    row_data.append(cell)
                else:
                    row_data.append(str(cell))
            data.append(row_data)
        
        # Extract headers (first row)
        headers = []
        if data:
            for i, val in enumerate(data[0]):
                headers.append({
                    "index": i,
                    "letter": get_column_letter(i),
                    "name": str(val) if val else f"Col{i+1}"
                })
        
        wb.close()
        
        return {
            "path": str(path.absolute()),
            "filename": path.name,
            "sheets": sheet_names,
            "activeSheet": sheet.title,
            "headers": headers,
            "data": data,
            "rowCount": len(data),
            "colCount": len(headers)
        }
        
    except ImportError:
        raise HTTPException(500, "openpyxl is not installed. Run: pip install openpyxl")
    except Exception as e:
        raise HTTPException(500, f"Error parsing Excel: {str(e)}")


@router.get("/browse")
async def browse_excel():
    """
    Open native file dialog to select Excel file.
    Uses tkinter for cross-platform dialog.
    """
    try:
        import tkinter as tk
        from tkinter import filedialog
        
        # Create hidden root window
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        root.focus_force()
        
        # Open file dialog
        filepath = filedialog.askopenfilename(
            title="Seleccionar archivo Excel",
            filetypes=[
                ("Excel files", "*.xlsx *.xls *.xlsm"),
                ("All files", "*.*")
            ]
        )
        
        root.destroy()
        
        if filepath:
            return {"path": filepath, "cancelled": False}
        return {"path": None, "cancelled": True}
        
    except Exception as e:
        raise HTTPException(500, f"Error opening file dialog: {str(e)}")


@router.post("/browse-and-parse")
async def browse_and_parse():
    """
    Combined: Open dialog, select file, parse and return data.
    Convenience endpoint for UI.
    """
    # First, browse
    browse_result = await browse_excel()
    
    if browse_result.get("cancelled") or not browse_result.get("path"):
        return {"cancelled": True}
    
    # Then parse
    parse_result = await parse_excel(ExcelParseRequest(path=browse_result["path"]))
    return parse_result


def get_column_letter(index: int) -> str:
    """Convert column index to Excel letter (0=A, 1=B, 26=AA, etc.)"""
    result = ""
    while index >= 0:
        result = chr(65 + index % 26) + result
        index = index // 26 - 1
    return result
