#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Project: DocuFlow
File:  run_validation_audit.py
Created: 2025-11-06
Author: @lewopxd

Description:
Script de "Auditoría". Lee el Excel maestro, determina qué actas
debería tener cada persona, valida si los archivos físicos existen,
y genera un reporte detallado en Excel con los hallazgos.

REGLA DE NEGOCIO: Una persona válida (que cumple CUALQUIER filtro)
debe tener EXACTAMENTE 3 actas. Ni más, ni menos.

MODIFICADO (v2):
- Se cambió la lógica de 'f.startswith(prefix)' a 'prefix in f'
  para que la auditoría funcione con los nombres de archivo
  que ahora tienen el prefijo de fecha (ej. "YYYYMMDD_ACTA...").
"""

import sys
import json
import re
import copy
import os
from pathlib import Path
from typing import Dict, Any, List

# --- [ INICIO: AJUSTE DE PYTHONPATH ] ---
try:
    current_dir = Path(__file__).parent
    project_root = current_dir.parent.parent 
    src_path = project_root / 'src'

    if str(src_path) not in sys.path:
        sys.path.insert(0, str(src_path))
        
    print(f"Ruta 'src' añadida a sys.path: {src_path}")
        
except NameError:
    print("Advertencia: No se pudo ajustar el sys.path automáticamente.")
    pass
# --- [ FIN: AJUSTE DE PYTHONPATH ] ---

try:
    from core.data_sheet.data_sheet_parser import get_row_data_sheet
    print("✓ Módulo 'get_row_data_sheet' importado.")
except ImportError as e:
    print(f"❌ CRITICAL: No se pudo importar 'get_row_data_sheet'.")
    print(f"   Asegúrate de que la carpeta 'src' está en el PYTHONPATH.")
    print(f"   Error: {e}")
    sys.exit(1)

# --- [ NUEVO IMPORT PARA EXCEL ] ---
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    print("✓ Módulo 'openpyxl' importado para generar reportes.")
except ImportError:
    print(f"❌ CRITICAL: Se necesita 'openpyxl' para crear el reporte Excel.")
    print(f"   Instálalo con: pip install openpyxl")
    sys.exit(1)


# -------------------------------------------------------------
# -------------------[   CONSTANTES DE AUDITORÍA   ]------------
# -------------------------------------------------------------

EXCEL_FILE_PATH = r"C:\Users\Admin\Desktop\FASE3-APROBADOS\DB-JOVENES_F3.xlsx"
OUTPUT_ALL_PATH = Path(r"C:\Users\Admin\Desktop\FASE3-APROBADOS\OUTPUT_ALL_BULK_2")
AUDIT_REPORT_NAME = "Auditoria_Actas_Generadas.xlsx"

# --- 1. Definir Hojas a Procesar ---
JOBS_TO_RUN = [
    {
        "name": "PRESENCIAL",
        "sheet": "PRESENCIAL",
        "table": "Table_1"
    },
    {
        "name": "VIRTUAL",
        "sheet": "VIRTUAL",
        "table": "Table_2"
    },
    {
        "name": "GUIAS",
        "sheet": "GUIAS",
        "table": "Table_3"
    }
]

# --- 2. Definir TODAS las columnas que se necesitarán ---
MASTER_COLUMN_LIST = [
    "NOMBRES", "APELLIDOS", "ESTADO INTERMEDIACION",
    "ACTUALIZACION", # Para Acta 1
    "E2. HABILIDADES BLANDAS", # Para Acta 2
    "E3. FERIA O AUTOPOSTULACION" # Para Acta 3
]

# --- 3. Definir TODAS las reglas de negocio y nombres de archivo ---
# NOTA: El "filename_prefix" ahora se buscará como "substring"
# (ej. "ACTA 1..." se encontrará en "20251022_ACTA 1...")
ACTA_DEFINITIONS = [
    {
        "name": "Acta 1 - Usuario Nuevo (APE)",
        "short_name": "A1 (Nuevo)",
        "filename_prefix": "ACTA 1 - APE - (USUARIO NUEVO)",
        "filter_rules": {
            "AND": [
                {"column": "ACTUALIZACION", "operator": "equal", "value": "NO"},
                {"column": "ESTADO INTERMEDIACION", "operator": "equal", "value": "OK"}
            ]
        }
    },
    {
        "name": "Acta 1 - Actualización (APE)",
        "short_name": "A1 (Actualiz.)",
        "filename_prefix": "ACTA 1 - APE - (ACTUALIZACION)",
        "filter_rules": {
            "AND": [
                {"column": "ACTUALIZACION", "operator": "equal", "value": "SI"},
                {"column": "ESTADO INTERMEDIACION", "operator": "equal", "value": "OK"}
            ]
        }
    },
    {
        "name": "Acta 2 - Hab. Presencial (PDF)",
        "short_name": "A2 (Presencial)",
        "filename_prefix": "ACTA 2 - PRESENCIAL",
        "filter_rules": {
            "AND": [
                {"column": "E2. HABILIDADES BLANDAS", "operator": "equal", "value": "Presencial"},
                {"column": "ESTADO INTERMEDIACION", "operator": "equal", "value": "OK"}
            ]
        }
    },
    {
        "name": "Acta 2 - Hab. Virtual (PDF)",
        "short_name": "A2 (Virtual)",
        "filename_prefix": "ACTA 2 - VIRTUAL",
        "filter_rules": {
            "AND": [
                {"column": "E2. HABILIDADES BLANDAS", "operator": "equal", "value": "Virtual"},
                {"column": "ESTADO INTERMEDIACION", "operator": "equal", "value": "OK"}
            ]
        }
    },
    {
        "name": "Acta 3 - Autopostulación",
        "short_name": "A3 (Autopost.)",
        "filename_prefix": "ACTA 3 - AUTOPOSTULACION",
        "filter_rules": {
            "AND": [
                {"column": "E3. FERIA O AUTOPOSTULACION", "operator": "equal", "value": "Autopostulación"},
                {"column": "ESTADO INTERMEDIACION", "operator": "equal", "value": "OK"}
            ]
        }
    },
    {
        "name": "Acta 3 - Feria ExpoEmpleo",
        "short_name": "A3 (Feria)",
        "filename_prefix": "ACTA 3 - FERIA EXPOEMPLEO OCT23",
        "filter_rules": {
            "AND": [
                {"column": "E3. FERIA O AUTOPOSTULACION", "operator": "equal", "value": "Feria"},
                {"column": "ESTADO INTERMEDIACION", "operator": "equal", "value": "OK"}
            ]
        }
    }
]

# -------------------------------------------------------------
# -------------------[   HELPER FUNCTIONS   ]------------------
# -------------------------------------------------------------

# --- [ Helpers de Lógica ] ---

def _evaluate_row(row_data: dict, rules: dict) -> bool:
    try:
        if "AND" in rules:
            return all(_evaluate_row(row_data, rule) for rule in rules["AND"])
        if "OR" in rules:
            return any(_evaluate_row(row_data, rule) for rule in rules["OR"])

        column = rules.get("column")
        op = rules.get("operator")
        value = rules.get("value")
        cell_value = row_data.get(column)
        
        if cell_value is not None:
             cell_value = str(cell_value)

        if op == "equal":
            return cell_value == value
        
        return False
    except Exception:
        return False

def _transform_trim(value: str) -> str:
    return value.strip()

def _transform_uppercase(value: str) -> str:
    return value.upper()

def _transform_regex_replace(value: str, pattern: str, replacement: str) -> str:
    try:
        return re.sub(pattern, replacement, value)
    except re.error:
        return value

def get_person_folder_name(row: dict) -> str:
    nombres = str(row.get("NOMBRES", ""))
    apellidos = str(row.get("APELLIDOS", ""))
    
    nombres = _transform_regex_replace(nombres, "\\b(NA|na)\\b", "")
    nombres = _transform_regex_replace(nombres, "\\s+", " ")
    nombres = _transform_trim(nombres)
    nombres = _transform_uppercase(nombres)
    
    apellidos = _transform_trim(apellidos)
    apellidos = _transform_uppercase(apellidos)
    
    folder_name = f"{nombres} {apellidos}".strip()
    
    invalid_chars = r'[*?:"<>|/\\]'
    folder_name = re.sub(invalid_chars, '_', folder_name)
    
    return folder_name

# --- [ Helpers de Estilo para Excel ] ---

FONT_BOLD = Font(bold=True)
FONT_BOLD_WHITE = Font(bold=True, color="FFFFFF")
FILL_HEADER = PatternFill(start_color="404040", end_color="404040", fill_type="solid")
FILL_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_GRAY = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

def setup_sheet_headers(ws, definitions):
    """Escribe y da estilo a las cabeceras de una hoja de auditoría."""
    headers = ["Fila Excel", "Nombre del Joven", "Carpeta Esperada"]
    acta_headers = [d["short_name"] for d in definitions]
    headers.extend(acta_headers)
    
    headers.append("Conteo de Actas") # Nueva columna
    headers.append("Estado Final")
    
    ws.append(headers)
    
    for cell in ws[1]: # Iterar la fila 1
        cell.font = FONT_BOLD_WHITE
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER

def style_cell(cell, status):
    """Aplica color a una celda basado en su estado."""
    cell.alignment = ALIGN_CENTER
    if status == "ENCONTRADO" or status == "COMPLETO":
        cell.fill = FILL_GREEN
    elif status == "FALTA" or status == "INCOMPLETO":
        cell.fill = FILL_RED
    elif status == "N/A":
        cell.fill = FILL_GRAY

def auto_fit_columns(ws):
    """Ajusta el ancho de todas las columnas en la hoja."""
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

# -------------------------------------------------------------
# -------------------[   EJECUCIÓN MAESTRA DE AUDITORÍA   ]----
# -------------------------------------------------------------

if __name__ == "__main__":
    print(f"╔══════════════════════════════════════════════════════════════╗")
    print(f"║               SCRIPT DE AUDITORÍA DE ACTAS (v2)              ║")
    print(f"╚══════════════════════════════════════════════════════════════╝")
    print(f"\nAuditando Directorio: {OUTPUT_ALL_PATH}")
    print(f"Contra Excel: {EXCEL_FILE_PATH}")
    
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    
    total_personas_auditadas = 0
    total_personas_con_error = 0 # Renombrado para claridad

    # 1. Iterar cada HOJA (PRESENCIAL, VIRTUAL, GUIAS)
    for job in JOBS_TO_RUN:
        sheet_name = job["sheet"]
        table_name = job["table"]
        job_name = job["name"]
        
        print(f"\n{'-'*70}")
        print(f"INICIANDO HOJA: \"{job_name}\" (Tabla: {table_name})")
        print(f"{'-'*70}")
        
        ws = wb.create_sheet(job_name)
        setup_sheet_headers(ws, ACTA_DEFINITIONS)
        current_excel_row = 2
        
        excel_config = {
            "filePath": EXCEL_FILE_PATH,
            "sheetName": sheet_name,
            "tableName": table_name,
            "rangeConfig": {"type": "all"},
            "columns": MASTER_COLUMN_LIST
        }
        
        row_data_list = get_row_data_sheet(excel_config)
        
        if not row_data_list:
            print(f"   ℹ️ No se encontraron filas en la hoja '{sheet_name}'. Saltando...")
            continue
            
        print(f"   ✓ Leídas {len(row_data_list)} filas de la hoja '{sheet_name}'. Auditando...")
        
        # 3. Iterar cada PERSONA (fila) en la hoja
        for row in row_data_list:
            row_index = row.get("row_index", "N/A")
            
            audit_results = {}
            expected_acta_prefixes = {}
            person_is_required = False
            
            # 4. Determinar qué actas DEBERÍA tener esta persona
            for definition in ACTA_DEFINITIONS:
                short_name = definition["short_name"]
                if _evaluate_row(row, definition["filter_rules"]):
                    person_is_required = True
                    expected_acta_prefixes[short_name] = definition["filename_prefix"]
                    audit_results[short_name] = "FALTA"
                else:
                    audit_results[short_name] = "N/A"
            
            if not person_is_required:
                continue
            
            # 5. Si la persona SÍ debe tener actas, iniciar reporte
            total_personas_auditadas += 1
            has_errors_this_row = False # Se basará en el conteo final
            
            person_folder_name = get_person_folder_name(row)
            person_full_name = f"{row.get('NOMBRES', '')} {row.get('APELLIDOS', '')}".strip()
            
            print(f"\n--- VALIDANDO: {person_full_name} (Fila: {row_index}, Carpeta: {person_folder_name})")
            
            # 6. Validar el sistema de archivos (la "Realidad")
            person_folder_path = OUTPUT_ALL_PATH / job_name / person_folder_name
            total_encontradas = 0
            
            if not person_folder_path.exists():
                print(f"   REALIDAD: ❌ ERROR: ¡La carpeta de la persona NO EXISTE!")
            else:
                try:
                    physical_files = os.listdir(person_folder_path)
                except Exception as e:
                    print(f"   REALIDAD: ❌ ERROR: No se pudo leer la carpeta. (Error: {e})")
                    physical_files = [] # Tratar como si no se hubiera encontrado nada

                # 7. Comparar Teoría vs Realidad
                for short_name, prefix in expected_acta_prefixes.items():
                    
                    # --- [ INICIO DE MODIFICACIÓN ] ---
                    # Antes: found = any(f.startswith(prefix) for f in physical_files)
                    # Ahora: Se busca el prefijo en CUALQUIER PARTE del nombre
                    found = any(prefix in f for f in physical_files)
                    # --- [ FIN DE MODIFICACIÓN ] ---
                    
                    if found:
                        print(f"       ✓ ENCONTRADA: {short_name}")
                        audit_results[short_name] = "ENCONTRADO"
                    else:
                        print(f"       ❌ FALTA: {short_name}")
                        # No se setea has_errors_this_row aquí
            
            # --- [ Lógica de Conteo ] ---
            
            # Calcular totales
            total_encontradas = list(audit_results.values()).count("ENCONTRADO")
            
            # Definir el estado basado en tu regla de "exactamente 3"
            conteo_status = f"{total_encontradas} / 3"
            
            if total_encontradas == 3:
                final_status = "COMPLETO"
                print(f"   --- RESULTADO: ✅ {conteo_status} - TODO CORRECTO ---")
            else:
                final_status = "INCOMPLETO"
                has_errors_this_row = True # Marcar la fila como errónea
                total_personas_con_error += 1
                print(f"   --- RESULTADO: ⚠️ {conteo_status} - CON ERRORES (Se esperaban 3) ---")

            # --- [ FIN Lógica de Conteo ] ---

            # 8. Escribir la fila completa en el Excel
            ws.cell(row=current_excel_row, column=1, value=row_index)
            ws.cell(row=current_excel_row, column=2, value=person_full_name)
            ws.cell(row=current_excel_row, column=3, value=person_folder_name)
            
            # Escribir las 6 columnas de estado de actas (diagnóstico)
            col_idx = 4
            for definition in ACTA_DEFINITIONS:
                status = audit_results[definition["short_name"]]
                cell = ws.cell(row=current_excel_row, column=col_idx, value=status)
                style_cell(cell, status)
                col_idx += 1
            
            # --- [ Escribir nuevas columnas ] ---
            
            # Columna "Conteo de Actas" (Col 10)
            cell_conteo = ws.cell(row=current_excel_row, column=col_idx, value=conteo_status)
            cell_conteo.font = FONT_BOLD
            style_cell(cell_conteo, final_status) # Colorear basado en el estado final
            col_idx += 1

            # Columna "Estado Final" (Col 11)
            cell_final = ws.cell(row=current_excel_row, column=col_idx, value=final_status)
            cell_final.font = FONT_BOLD
            style_cell(cell_final, final_status)
            
            # --- [ FIN Escribir nuevas columnas ] ---
            
            current_excel_row += 1
        
        auto_fit_columns(ws)

    # 9. Reporte Final en Consola
    print(f"\n{'-'*70}")
    print("AUDITORÍA FINALIZADA.")
    print(f"📊 RESUMEN TOTAL DE AUDITORÍA:")
    print(f"   - Personas auditadas (que requerían actas): {total_personas_auditadas}")
    print(f"   - Personas con errores (conteo != 3): {total_personas_con_error}") # Modificado
    
    # 10. Guardar el archivo Excel
    try:
        report_path = OUTPUT_ALL_PATH / AUDIT_REPORT_NAME
        wb.save(report_path)
        print(f"\n✅ ¡Reporte de auditoría guardado exitosamente en:")
        print(f"   {report_path}")
    except Exception as e:
        print(f"\n❌ ERROR CRÍTICO AL GUARDAR EL REPORTE:")
        print(f"   {e}")
        print(f"   Asegúrate de que el archivo '{report_path}' no esté abierto.")
        
    if total_personas_con_error == 0 and total_personas_auditadas > 0:
        print(f"\n✅ ¡FELICITACIONES! Todas las personas auditadas tienen 3 actas.")
    elif total_personas_auditadas == 0:
        print(f"\nℹ️ No se encontró ninguna persona que cumpliera con los criterios de filtro.")
    else:
        print(f"\n⚠️ ¡ATENCIÓN! Se encontraron {total_personas_con_error} personas con un conteo incorrecto de actas.")
        
    print(f"\n╔══════════════════════════════════════════════════════════════╗")
    print(f"║                     PROCESO FINALIZADO                       ║")
    print(f"╚══════════════════════════════════════════════════════════════╝")