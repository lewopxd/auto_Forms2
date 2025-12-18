#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Project: DocuFlow Converter PRO
File: run_converter_pro.py
Author: @lewopxd
Created: 2025-12-07

Description:
Script de alto rendimiento para conversión masiva Word -> PDF.
Incluye selectores de carpeta, normalización de nombres,
mapeo de renombramiento y reportes detallados en Excel.
Requires: pip install docx2pdf openpyxl
"""

import sys
import os
import shutil
import unicodedata
import logging
from pathlib import Path
from datetime import datetime

# --- [ CONFIGURACIÓN DE LIBRERÍAS ] ---
try:
    from docx2pdf import convert
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError as e:
    print(f"❌ ERROR CRÍTICO: Falta una librería necesaria: {e}")
    print("   Ejecuta: pip install docx2pdf openpyxl")
    sys.exit(1)

# =============================================================================
# ⚙️ CONFIGURACIÓN MAESTRA
# =============================================================================

CONFIG = {
    # 📂 Ruta Principal (Donde están las carpetas de los jóvenes)
    "BASE_PATH": Path(r"C:\Users\Admin\Desktop\PLAYGROUND MARCELA\F2 CURSOS CORTOS\OUTPUT\CERTIFICADOS"),

    # 🎯 Configuración de Ejecución
    # MODE: "ALL" (Recorre todo) | "SELECT" (Busca carpeta específica)
    # KEY: Nombre del joven a buscar (Solo si MODE="SELECT")
    "EXECUTION": {
        "MODE": "SELECT",       
        "KEY": "ADRIANA MORENO JIMENEZ"  
    },

    # 📄 Mapeo de Conversión y Renombramiento
    # Si el nombre del archivo contiene la KEY, se convierte y se renombra al VALUE.
    "FILE_MAP": {
        "ACTA 1": "001ActaInicioRuta.pdf",
        "ACTA 2": "002ActaPermanenciaRuta.pdf",
        "ACTA 3": "003ActaFinalizacionRuta.pdf"
    },

    # 💾 Configuración de Salida
    # MODE: 
    #   "SAME_PLACE" -> Guarda junto al Word.
    #   "SUBFOLDER"  -> Crea una carpeta dentro de la del joven (nombre en PATH).
    #   "SPECIFIC"   -> Una ruta absoluta única para todos los archivos.
    "OUTPUT": {
        "MODE": "SAME_PLACE",
        "PATH": "PDF_FINALES"  # Solo se usa si MODE es SUBFOLDER o SPECIFIC
    },

    # 🛡️ Políticas de Archivo
    "POLICIES": {
        "OVERWRITE": True,      # ¿Sobrescribir si el PDF ya existe?
        "STOP_ON_ERROR": False  # ¿Detener todo si falla un archivo?
    }
}

# =============================================================================
# 🛠️ UTILIDADES (HELPERS)
# =============================================================================

def setup_logger():
    """Configura el logger para consola."""
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        datefmt='%H:%M:%S'
    )
    return logging.getLogger("DocuFlowConverter")

logger = setup_logger()

def normalize_text(text: str) -> str:
    """
    Normaliza texto para búsquedas seguras:
    - Minúsculas
    - Sin tildes
    - Sin espacios extra
    """
    if not text: return ""
    s = str(text).strip().lower()
    s = unicodedata.normalize('NFKD', s).encode('ASCII', 'ignore').decode('ASCII')
    return " ".join(s.split())

def determine_output_folder(source_folder: Path) -> Path:
    """Calcula la carpeta de destino según la configuración."""
    mode = CONFIG["OUTPUT"]["MODE"]
    path_val = CONFIG["OUTPUT"]["PATH"]

    if mode == "SAME_PLACE":
        return source_folder
    elif mode == "SUBFOLDER":
        target = source_folder / path_val
        target.mkdir(parents=True, exist_ok=True)
        return target
    elif mode == "SPECIFIC":
        target = Path(path_val)
        target.mkdir(parents=True, exist_ok=True)
        return target
    else:
        logger.warning(f"Modo de salida '{mode}' desconocido. Usando SAME_PLACE.")
        return source_folder

def get_target_folders(base_path: Path) -> list:
    """Identifica qué carpetas procesar según el modo ALL o SELECT."""
    if not base_path.exists():
        logger.error(f"La ruta base no existe: {base_path}")
        return []

    all_subfolders = [f for f in base_path.iterdir() if f.is_dir()]
    
    if CONFIG["EXECUTION"]["MODE"] == "ALL":
        return all_subfolders
    
    elif CONFIG["EXECUTION"]["MODE"] == "SELECT":
        search_key = normalize_text(CONFIG["EXECUTION"]["KEY"])
        logger.info(f"🔍 Modo SELECT activado. Buscando: '{search_key}'")
        
        matches = []
        for folder in all_subfolders:
            if search_key in normalize_text(folder.name):
                matches.append(folder)
        
        if not matches:
            logger.warning("⚠️ No se encontraron carpetas que coincidan con la búsqueda.")
        return matches
    
    return []

# =============================================================================
# 🚀 NÚCLEO DE CONVERSIÓN
# =============================================================================

def process_conversion():
    logger.info("🚀 INICIANDO PROCESO DE CONVERSIÓN MASIVA")
    logger.info(f"📂 Ruta Base: {CONFIG['BASE_PATH']}")
    
    report_data = []
    folders_to_process = get_target_folders(CONFIG["BASE_PATH"])
    
    logger.info(f"📋 Carpetas a procesar: {len(folders_to_process)}")

    for folder in folders_to_process:
        logger.info(f"🔸 Procesando carpeta: {folder.name}")
        
        # Determinar dónde guardar los PDFs de esta carpeta
        target_output_folder = determine_output_folder(folder)

        # Listar archivos Word
        word_files = list(folder.glob("*.docx"))
        
        if not word_files:
            logger.info("   ⚠️ No hay archivos .docx en esta carpeta.")
            continue

        for word_file in word_files:
            file_name_upper = word_file.name.upper()
            processed = False

            # Verificar contra el Mapa de Archivos
            for key, final_name_pdf in CONFIG["FILE_MAP"].items():
                if key.upper() in file_name_upper:
                    # ¡Match encontrado!
                    target_pdf_path = target_output_folder / final_name_pdf
                    
                    # Chequeo de política de sobrescritura
                    if target_pdf_path.exists() and not CONFIG["POLICIES"]["OVERWRITE"]:
                        status = "OMITIDO (Existe)"
                        logger.info(f"   ⏭️ {status}: {target_pdf_path.name}")
                        report_data.append({
                            "Carpeta": folder.name,
                            "Archivo Origen": word_file.name,
                            "Archivo Destino": final_name_pdf,
                            "Estado": status,
                            "Detalle": "El archivo ya existe y OVERWRITE=False"
                        })
                        processed = True
                        break # Salir del loop de keys, ya encontramos match

                    # Ejecutar Conversión
                    try:
                        logger.info(f"   🔄 Convirtiendo: {word_file.name} -> {final_name_pdf}")
                        
                        # docx2pdf convert(input, output)
                        convert(str(word_file), str(target_pdf_path))
                        
                        if target_pdf_path.exists():
                            status = "EXITO"
                            logger.info(f"   ✅ Generado: {target_pdf_path.name}")
                        else:
                            status = "FALLO (No generado)"
                            logger.error(f"   ❌ El archivo no apareció tras conversión.")

                        report_data.append({
                            "Carpeta": folder.name,
                            "Archivo Origen": word_file.name,
                            "Archivo Destino": final_name_pdf,
                            "Estado": status,
                            "Detalle": "Conversión completada" if status == "EXITO" else "docx2pdf no generó salida"
                        })

                    except Exception as e:
                        status = "ERROR CRITICO"
                        err_msg = str(e)
                        logger.error(f"   ❌ Error convirtiendo {word_file.name}: {err_msg}")
                        report_data.append({
                            "Carpeta": folder.name,
                            "Archivo Origen": word_file.name,
                            "Archivo Destino": final_name_pdf,
                            "Estado": status,
                            "Detalle": err_msg
                        })
                        if CONFIG["POLICIES"]["STOP_ON_ERROR"]:
                            logger.critical("🛑 Deteniendo ejecución por política STOP_ON_ERROR")
                            generate_excel_report(report_data)
                            sys.exit(1)

                    processed = True
                    break # Salir del loop de keys (match único)
            
            # Si el archivo no coincidió con ninguna key
            if not processed:
                # Opcional: Reportar archivos ignorados
                pass

    generate_excel_report(report_data)

# =============================================================================
# 📊 GENERADOR DE REPORTE
# =============================================================================

def generate_excel_report(data):
    if not data:
        logger.warning("⚠️ No se generaron datos para el reporte.")
        return

    logger.info("📊 Generando reporte Excel...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Conversión"

    # Encabezados
    headers = ["Carpeta Joven", "Archivo Word (Origen)", "PDF Generado (Destino)", "Estado", "Detalles"]
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Datos
    for r, item in enumerate(data, 2):
        ws.cell(row=r, column=1, value=item["Carpeta"])
        ws.cell(row=r, column=2, value=item["Archivo Origen"])
        ws.cell(row=r, column=3, value=item["Archivo Destino"])
        
        status_cell = ws.cell(row=r, column=4, value=item["Estado"])
        # Color condicional
        if "EXITO" in item["Estado"]:
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Verde
        elif "ERROR" in item["Estado"] or "FALLO" in item["Estado"]:
            status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Rojo
        else:
            status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid") # Amarillo

        ws.cell(row=r, column=5, value=item["Detalle"])

    # Ajustar ancho columnas
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = min(adjusted_width, 50) # Max ancho 50

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = CONFIG["BASE_PATH"] / f"Reporte_Conversion_{timestamp}.xlsx"
    
    try:
        wb.save(report_path)
        logger.info(f"✅ Reporte guardado en: {report_path}")
    except Exception as e:
        logger.error(f"❌ No se pudo guardar el Excel (¿quizás está abierto?): {e}")

# =============================================================================
# 🏁 MAIN
# =============================================================================

if __name__ == "__main__":
    print(f"\n╔══════════════════════════════════════════════════════════════╗")
    print(f"║   CONVERTIDOR PRO WORD -> PDF (DocuFlow Standalone)          ║")
    print(f"╚══════════════════════════════════════════════════════════════╝")
    
    if not CONFIG["BASE_PATH"].exists():
        logger.error(f"La ruta base no existe: {CONFIG['BASE_PATH']}")
        sys.exit(1)

    process_conversion()
    
    print("\n✅ Proceso finalizado.")